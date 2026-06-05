from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "docs" / "CargoPilot-Integration-Partner-Guide.docx"
WORK_DIR = ROOT / ".docx-work" / "integration-partner-guide"

NAVY = "0B2545"
DEEP_BLUE = "163A5F"
TEAL = "078A83"
CYAN = "147D92"
INK = "172033"
MUTED = "5B677A"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E7F6F4"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF5D6"
GOLD = "8B6508"
PALE_RED = "FDECEC"
RED = "9B1C1C"
WHITE = "FFFFFF"
BORDER = "CDD6E1"
CODE_BG = "F4F6F9"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

FONT_NAME = "Calibri"
MONO_FONT = "Consolas"
WIN_FONT = Path("C:/Windows/Fonts/segoeui.ttf")
WIN_FONT_BOLD = Path("C:/Windows/Fonts/segoeuib.ttf")


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name: str = FONT_NAME, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def clear_paragraph_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for style_name, size, color, before, after in (
        ("Heading 1", 16, DEEP_BLUE, 18, 10),
        ("Heading 2", 13, CYAN, 14, 7),
        ("Heading 3", 11.5, DEEP_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style.font.size = Pt(10.5)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.2

    caption = styles["Caption"]
    caption.font.name = FONT_NAME
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)


def configure_header_footer(section, first_page: bool = False) -> None:
    section.different_first_page_header_footer = first_page
    header = section.header
    p = header.paragraphs[0]
    clear_paragraph_content(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("CargoPilot  |  Partner Integration Specification")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    clear_paragraph_content(p)
    add_page_number(p)


def add_title(doc: Document, text: str, size: float = 28, color: str = NAVY,
              alignment=WD_ALIGN_PARAGRAPH.LEFT, after: float = 8) -> None:
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=True)


def add_subtitle(doc: Document, text: str, size: float = 13, color: str = MUTED,
                 alignment=WD_ALIGN_PARAGRAPH.LEFT, after: float = 16) -> None:
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color)


def add_kicker(doc: Document, text: str, color: str = TEAL, after: float = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text.upper())
    set_run_font(run, size=9, color=color, bold=True)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, color=INK)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        run = p.add_run(item)
        set_run_font(run, color=INK)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        run = p.add_run(item)
        set_run_font(run, color=INK)


def add_callout(doc: Document, label: str, text: str, fill: str = LIGHT_TEAL,
                accent: str = TEAL) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA], 0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, accent, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA], 0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_BG)
    set_cell_border(cell, BORDER, "4")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(code.strip("\n").splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name=MONO_FONT, size=8.2, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]],
              widths_dxa: Sequence[int], header_fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths_dxa)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, color=NAVY, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_border(cell)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cell, "FAFBFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=9, color=INK)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.45) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def font(size: int, bold: bool = False):
    path = WIN_FONT_BOLD if bold else WIN_FONT
    if path.exists():
        return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def rounded_box(draw: ImageDraw.ImageDraw, xy, fill, outline, title, subtitle=None,
                title_color=WHITE, subtitle_color=WHITE, radius=20) -> None:
    if not str(title_color).startswith("#") and str(title_color).lower() not in {"white", "black"}:
        title_color = f"#{title_color}"
    if not str(subtitle_color).startswith("#") and str(subtitle_color).lower() not in {"white", "black"}:
        subtitle_color = f"#{subtitle_color}"
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    cx = (x1 + x2) // 2
    title_font = font(30, True)
    sub_font = font(21)
    title_bbox = draw.textbbox((0, 0), title, font=title_font)
    title_y = (y1 + y2) // 2 - (title_bbox[3] - title_bbox[1]) // 2
    if subtitle:
        title_y -= 22
    draw.text((cx, title_y), title, font=title_font, fill=title_color, anchor="mm")
    if subtitle:
        draw.text((cx, title_y + 43), subtitle, font=sub_font, fill=subtitle_color, anchor="mm")


def arrow(draw: ImageDraw.ImageDraw, start, end, color=DEEP_BLUE, width=6) -> None:
    draw.line([start, end], fill=f"#{color}", width=width)
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    size = 18
    p1 = (x2, y2)
    p2 = (x2 - ux * size + px * size * 0.55, y2 - uy * size + py * size * 0.55)
    p3 = (x2 - ux * size - px * size * 0.55, y2 - uy * size - py * size * 0.55)
    draw.polygon([p1, p2, p3], fill=f"#{color}")


def create_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 1050), "white")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((60, 90, 1230, 970), radius=30, fill="#F4F7FA", outline="#CDD6E1", width=4)
    d.text((95, 115), "CARGOPILOT ERP", font=font(30, True), fill=f"#{NAVY}")
    rounded_box(d, (140, 240, 500, 400), f"#{NAVY}", f"#{NAVY}", "Core Modules", "Orders | Tracking | Support")
    rounded_box(d, (650, 240, 1010, 400), f"#{DEEP_BLUE}", f"#{DEEP_BLUE}", "Integration Outbox", "Durable event queue")
    rounded_box(d, (650, 520, 1010, 680), f"#{TEAL}", f"#{TEAL}", "Integration Worker", "Dispatch | Retry | DLQ")
    rounded_box(d, (140, 520, 500, 680), "#FFFFFF", f"#{CYAN}", "Provider Registry", "Company | Domain | Environment", NAVY, MUTED)
    rounded_box(d, (140, 760, 500, 900), "#FFFFFF", f"#{CYAN}", "Encrypted Secrets", "URL | Token | API key", NAVY, MUTED)
    rounded_box(d, (650, 760, 1010, 900), "#FFFFFF", f"#{CYAN}", "Webhook Gateway", "Verify | Deduplicate | Normalize", NAVY, MUTED)
    d.rounded_rectangle((1300, 90, 1740, 970), radius=30, fill="#F0FBFA", outline="#8BD0CA", width=4)
    d.text((1340, 115), "EXTERNAL PARTNERS", font=font(30, True), fill=f"#{NAVY}")
    rounded_box(d, (1360, 240, 1680, 380), f"#{TEAL}", f"#{TEAL}", "Carrier API")
    rounded_box(d, (1360, 455, 1680, 595), f"#{CYAN}", f"#{CYAN}", "SMS API")
    rounded_box(d, (1360, 670, 1680, 810), f"#{DEEP_BLUE}", f"#{DEEP_BLUE}", "Webhook Endpoint")
    arrow(d, (500, 320), (650, 320))
    arrow(d, (830, 400), (830, 520))
    arrow(d, (1010, 590), (1360, 310))
    arrow(d, (1010, 600), (1360, 525))
    arrow(d, (1010, 620), (1360, 740))
    arrow(d, (650, 830), (500, 830))
    arrow(d, (1360, 790), (1010, 830))
    arrow(d, (500, 590), (650, 590))
    arrow(d, (500, 830), (650, 650))
    img.save(path)


def create_sequence_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 1120), "white")
    d = ImageDraw.Draw(img)
    actors = [("Core", 180), ("Outbox", 500), ("Worker", 820), ("Registry", 1140), ("Partner API", 1510)]
    for label, x in actors:
        rounded_box(d, (x - 130, 50, x + 130, 150), f"#{NAVY if label == 'Core' else TEAL if label == 'Partner API' else DEEP_BLUE}", f"#{NAVY}", label)
        d.line([(x, 150), (x, 1030)], fill="#A9B7C6", width=3)
    steps = [
        (210, 180, 500, "1  Enqueue canonical event"),
        (310, 820, 500, "2  Claim pending record"),
        (410, 820, 1140, "3  Resolve active provider"),
        (510, 820, 1510, "4  Send request + idempotency key"),
        (620, 1510, 820, "5  Return response"),
        (730, 820, 500, "6  Persist attempt result"),
        (850, 820, 500, "7  Mark sent / retry / DLQ"),
    ]
    for y, x1, x2, label in steps:
        arrow(d, (x1, y), (x2, y), TEAL if x1 < x2 else DEEP_BLUE, 5)
        d.text(((x1 + x2) // 2, y - 30), label, font=font(20, True), fill=f"#{INK}", anchor="mm")
    d.rounded_rectangle((620, 930, 1020, 1020), radius=18, fill="#FFF5D6", outline="#C79A28", width=3)
    d.text((820, 975), "Temporary errors are retried.\nFinal failures move to dead letter.", font=font(20, True), fill=f"#{GOLD}", anchor="mm", align="center")
    img.save(path)


def create_state_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 930), "white")
    d = ImageDraw.Draw(img)
    boxes = {
        "pending": (170, 200, 470, 340, DEEP_BLUE),
        "processing": (650, 200, 1000, 340, TEAL),
        "sent": (1280, 100, 1580, 240, "177245"),
        "failed": (650, 520, 1000, 660, GOLD),
        "dead_letter": (1280, 520, 1580, 660, RED),
    }
    for label, (x1, y1, x2, y2, color) in boxes.items():
        rounded_box(d, (x1, y1, x2, y2), f"#{color}", f"#{color}", label.replace("_", " ").title())
    arrow(d, (470, 270), (650, 270), DEEP_BLUE)
    d.text((560, 235), "worker claims", font=font(19, True), fill=f"#{MUTED}", anchor="mm")
    arrow(d, (1000, 240), (1280, 170), "177245")
    d.text((1140, 175), "2xx", font=font(19, True), fill="#177245", anchor="mm")
    arrow(d, (820, 340), (820, 520), GOLD)
    d.text((900, 430), "retryable failure", font=font(19, True), fill=f"#{GOLD}", anchor="mm")
    arrow(d, (1000, 590), (1280, 590), RED)
    d.text((1140, 555), "attempts exhausted", font=font(19, True), fill=f"#{RED}", anchor="mm")
    arrow(d, (650, 590), (470, 330), TEAL)
    d.text((505, 500), "retry due", font=font(19, True), fill=f"#{TEAL}", anchor="mm")
    arrow(d, (1430, 520), (470, 310), TEAL)
    d.text((1190, 405), "admin replay creates new pending record", font=font(19, True), fill=f"#{TEAL}", anchor="mm")
    img.save(path)


def create_class_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 1160), "white")
    d = ImageDraw.Draw(img)

    def class_box(x, y, w, title, fields, color):
        height = 95 + len(fields) * 32
        d.rounded_rectangle((x, y, x + w, y + height), radius=18, fill="white", outline=f"#{color}", width=4)
        d.rounded_rectangle((x, y, x + w, y + 62), radius=18, fill=f"#{color}", outline=f"#{color}", width=4)
        d.rectangle((x, y + 42, x + w, y + 62), fill=f"#{color}")
        d.text((x + w // 2, y + 31), title, font=font(24, True), fill="white", anchor="mm")
        for idx, field in enumerate(fields):
            d.text((x + 22, y + 84 + idx * 32), field, font=font(18), fill=f"#{INK}")
        return (x, y, x + w, y + height)

    provider = class_box(90, 80, 610, "IntegrationProvider", [
        "id : UUID", "companyId : UUID", "domain : IntegrationDomain",
        "providerCode : string", "status : ProviderStatus", "environment : Environment"
    ], NAVY)
    secret = class_box(90, 620, 610, "IntegrationProviderSecret", [
        "id : UUID", "providerId : UUID", "keyVersion : integer",
        "encryptedSecretJson : string", "rotatedAt : datetime"
    ], CYAN)
    outbox = class_box(1050, 80, 650, "IntegrationOutbox", [
        "id : UUID", "companyId : UUID", "providerId : UUID?", "eventType : string",
        "status : OutboxStatus", "attemptCount : integer", "idempotencyKey : string", "payload : JSON"
    ], DEEP_BLUE)
    attempt = class_box(1050, 670, 650, "IntegrationDeliveryAttempt", [
        "id : UUID", "outboxId : UUID", "attemptNo : integer",
        "outcome : string", "statusCode : integer?", "retryable : boolean", "errorMessage : string?"
    ], TEAL)
    arrow(d, (395, provider[3]), (395, secret[1]), CYAN)
    d.text((430, 520), "1 to many", font=font(20, True), fill=f"#{MUTED}")
    arrow(d, (provider[2], 260), (outbox[0], 260), DEEP_BLUE)
    d.text((875, 220), "1 to many", font=font(20, True), fill=f"#{MUTED}", anchor="mm")
    arrow(d, (1375, outbox[3]), (1375, attempt[1]), TEAL)
    d.text((1410, 600), "1 to many", font=font(20, True), fill=f"#{MUTED}")
    img.save(path)


def create_diagrams() -> dict[str, Path]:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    paths = {
        "architecture": WORK_DIR / "architecture.png",
        "sequence": WORK_DIR / "outbound-sequence.png",
        "state": WORK_DIR / "outbox-state.png",
        "class": WORK_DIR / "integration-class.png",
    }
    create_architecture_diagram(paths["architecture"])
    create_sequence_diagram(paths["sequence"])
    create_state_diagram(paths["state"])
    create_class_diagram(paths["class"])
    return paths


def build_document() -> None:
    diagrams = create_diagrams()
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    section = doc.sections[0]
    configure_header_footer(section, first_page=True)

    # Cover page: customer_pack pattern, compact reference guide preset.
    doc.add_paragraph().paragraph_format.space_after = Pt(28)
    add_kicker(doc, "Partner Integration Specification", TEAL, 6)
    add_title(doc, "CargoPilot Integration Partner Guide", 30, NAVY, WD_ALIGN_PARAGRAPH.LEFT, 8)
    add_subtitle(
        doc,
        "Secure, reliable integration contracts for carriers, SMS providers, webhook partners, and enterprise systems.",
        14,
        MUTED,
        WD_ALIGN_PARAGRAPH.LEFT,
        24,
    )

    add_table(
        doc,
        ["Document", "Value"],
        [
            ["Version", "1.0"],
            ["Published", "June 3, 2026"],
            ["Audience", "External partners and implementation teams"],
            ["Classification", "Partner use"],
            ["Support", "Assigned during partner onboarding"],
        ],
        [2500, 6860],
        LIGHT_TEAL,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    add_callout(
        doc,
        "Integration promise",
        "CargoPilot keeps core ERP domains provider-agnostic while providing encrypted credentials, durable delivery, secure webhooks, idempotency, retries, and operational auditability.",
        LIGHT_BLUE,
        DEEP_BLUE,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("CargoPilot")
    set_run_font(r, size=19, color=TEAL, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Integration Layer | External Partner Documentation")
    set_run_font(r2, size=10, color=MUTED)
    doc.add_page_break()

    add_kicker(doc, "Document Navigation", TEAL)
    add_title(doc, "Contents", 24, NAVY, after=14)
    contents = [
        "1. Purpose and scope",
        "2. Integration operating model",
        "3. Partner onboarding",
        "4. Security, authentication, and idempotency",
        "5. Generic carrier API contract",
        "6. Generic SMS API contract",
        "7. Generic outbound webhook contract",
        "8. Inbound webhook contract",
        "9. Reliability and operations",
        "10. UML data model",
        "11. Testing and production readiness",
        "Appendices: events, errors, and reference examples",
    ]
    add_numbered(doc, contents)
    add_callout(
        doc,
        "How to use this guide",
        "Partners that can implement the generic contract can integrate without a custom adapter. Partners with provider-specific protocols should use this guide as the canonical business and security specification while CargoPilot implements a dedicated adapter.",
        LIGHT_TEAL,
        TEAL,
    )
    doc.add_page_break()

    doc.add_heading("1. Purpose and Scope", level=1)
    add_body(
        doc,
        "This guide defines how external companies connect to CargoPilot ERP. It covers outbound carrier, SMS, and webhook delivery; inbound signed webhooks; provider configuration; idempotency; retries; dead-letter handling; testing; and production onboarding.",
    )
    add_body(
        doc,
        "CargoPilot integrations are company-scoped. A provider configuration belongs to one company, one integration domain, and one environment. This allows different companies to use different partners, credentials, limits, and activation states.",
    )
    doc.add_heading("Integration domains", level=2)
    add_table(
        doc,
        ["Domain", "Purpose", "Current generic contract"],
        [
            ["carrier", "Create, cancel, and track partner shipments", "Available"],
            ["sms", "Send SMS and query delivery status", "Available"],
            ["webhook_sink", "Deliver CargoPilot events to partner endpoints", "Available"],
            ["payment", "Reserved for payment-provider convergence", "Provider-specific today"],
        ],
        [1700, 4200, 3460],
    )
    add_callout(
        doc,
        "API base URL",
        "Sandbox and production CargoPilot base URLs are supplied during onboarding. All production communication must use HTTPS.",
        PALE_GOLD,
        GOLD,
    )

    doc.add_heading("2. Integration Operating Model", level=1)
    add_body(
        doc,
        "CargoPilot core modules do not call partner APIs directly. Core modules emit canonical events into a durable integration outbox. A worker resolves the active provider, decrypts the provider configuration, dispatches the request, and records every attempt.",
    )
    add_figure(doc, diagrams["architecture"], "Figure 1. CargoPilot integration layer and external partner boundaries.")
    doc.add_heading("Core components", level=2)
    add_table(
        doc,
        ["Component", "Responsibility"],
        [
            ["Provider Registry", "Stores company, domain, provider code, environment, status, capabilities, timeout, and rate limit."],
            ["Encrypted Secret Store", "Stores rotatable provider endpoint and credential configuration encrypted at rest."],
            ["Integration Outbox", "Durably records outbound integration events before delivery."],
            ["Integration Worker", "Claims records, dispatches requests, handles retries, and moves final failures to dead letter."],
            ["Delivery Attempts", "Records each request result, provider status code, retryability, and error."],
            ["Webhook Gateway", "Verifies signatures, rejects stale or invalid requests, deduplicates events, and normalizes payloads."],
            ["Admin Console", "Allows authorized administrators to configure providers and operate the delivery queue."],
        ],
        [2450, 6910],
    )
    doc.add_heading("Provider lifecycle", level=2)
    add_table(
        doc,
        ["Status", "Meaning", "Dispatch behavior"],
        [
            ["active", "Configured and approved for use", "Worker can dispatch"],
            ["paused", "Temporarily suspended", "Worker does not dispatch"],
            ["disabled", "Not available for use", "Worker does not dispatch"],
        ],
        [1800, 3600, 3960],
    )

    doc.add_heading("3. Partner Onboarding", level=1)
    add_body(doc, "Every partner integration follows a controlled onboarding process.")
    add_numbered(
        doc,
        [
            "Agree on the integration domain and supported business operations.",
            "Confirm whether the generic CargoPilot contract is sufficient or a custom adapter is required.",
            "Exchange sandbox URLs, credentials, webhook secrets, and technical contacts.",
            "Configure the provider in CargoPilot using the sandbox environment.",
            "Verify outbound requests, response mapping, idempotency, and temporary failure behavior.",
            "Verify inbound webhook signature generation, event IDs, and duplicate handling.",
            "Complete production readiness checks and load production credentials.",
            "Activate the production provider after joint approval.",
        ],
    )
    doc.add_heading("When a custom adapter is required", level=2)
    add_bullets(
        doc,
        [
            "The partner cannot expose the generic `/shipments` or `/messages` endpoints.",
            "Authentication requires OAuth refresh, mutual TLS, or a provider-specific signature scheme.",
            "The partner uses XML, SOAP, multipart files, or multi-step workflows.",
            "Shipment labels, manifests, tracking statuses, or errors require provider-specific mapping.",
            "The partner contract has special rate limits, reconciliation, or asynchronous confirmation rules.",
        ],
    )
    add_callout(
        doc,
        "Architecture rule",
        "A custom adapter belongs inside CargoPilot integrations-core. Orders, payments, pricing, tracking, and support remain provider-agnostic.",
        LIGHT_BLUE,
        DEEP_BLUE,
    )

    doc.add_heading("4. Security, Authentication, and Idempotency", level=1)
    doc.add_heading("Outbound authentication", level=2)
    add_table(
        doc,
        ["Secret field", "Behavior"],
        [
            ["baseUrl", "Base URL used by carrier and SMS adapters."],
            ["endpointUrl", "Full target URL used by generic outbound webhook delivery."],
            ["token", "Sent as `Authorization: Bearer <token>`."],
            ["apiKey", "Sent in the configured API key header."],
            ["apiKeyHeader", "Header name for `apiKey`; defaults to `x-api-key`."],
        ],
        [2200, 7160],
    )
    add_code_block(
        doc,
        """
{
  "baseUrl": "https://partner.example.com",
  "token": "optional-bearer-token",
  "apiKey": "optional-api-key",
  "apiKeyHeader": "x-api-key"
}
""",
    )
    doc.add_heading("Idempotency requirements", level=2)
    add_bullets(
        doc,
        [
            "CargoPilot sends `x-idempotency-key` for outbound business operations.",
            "A repeated idempotency key must not create duplicate shipments, SMS messages, invoices, or payments.",
            "If the first request succeeded but its response was lost, a repeat request should return the original business reference.",
            "Inbound webhook `eventId` values must be stable and unique per provider.",
        ],
    )
    add_callout(
        doc,
        "Security requirement",
        "Partners must never place passwords, card data, personal documents, or long-lived secrets in URLs or webhook payloads.",
        PALE_RED,
        RED,
    )
    doc.add_heading("Headers sent by CargoPilot", level=2)
    add_code_block(
        doc,
        """
content-type: application/json
x-cargopilot-request-id: <outbox-record-id>
x-cargopilot-company-id: <company-id>
x-idempotency-key: <stable-idempotency-key>
authorization: Bearer <token>        # when configured
x-api-key: <api-key>                 # when configured
""",
    )

    doc.add_heading("5. Generic Carrier API Contract", level=1)
    add_body(
        doc,
        "A carrier that supports the following endpoints can integrate through the generic HTTP carrier adapter. All responses must be JSON.",
    )
    doc.add_heading("5.1 Create shipment", level=2)
    add_code_block(doc, "POST /shipments")
    add_code_block(
        doc,
        """
{
  "externalOrderId": "990000000123",
  "sender": {
    "name": "Sender Name",
    "phone": "+998901234567",
    "address": "Tashkent, Uzbekistan",
    "lat": 41.311081,
    "lng": 69.240562
  },
  "receiver": {
    "name": "Receiver Name",
    "phone": "+998909876543",
    "address": "Almaty, Kazakhstan",
    "lat": 43.238949,
    "lng": 76.889709
  },
  "parcels": [
    { "weightKg": 2.5, "quantity": 1, "description": "Documents" }
  ],
  "declaredValueMinor": "1200000",
  "currency": "UZS",
  "transportMode": "air",
  "serviceCode": "express",
  "metadata": { "source": "cargopilot" }
}
""",
    )
    add_body(doc, "Successful response:")
    add_code_block(
        doc,
        """
{
  "partnerShipmentId": "DHL-123456789",
  "trackingNumber": "JD014600011234567890",
  "labelUrl": "https://partner.example.com/labels/JD014600011234567890.pdf"
}
""",
    )
    add_callout(
        doc,
        "Required response",
        "`partnerShipmentId` is required. CargoPilot also accepts `shipmentId` or `id` as aliases.",
        LIGHT_TEAL,
        TEAL,
    )
    doc.add_heading("5.2 Cancel shipment", level=2)
    add_code_block(
        doc,
        """
POST /shipments/{partnerShipmentId}/cancel

{ "reason": "Customer requested cancellation" }
""",
    )
    doc.add_heading("5.3 Track shipment", level=2)
    add_code_block(
        doc,
        """
GET /shipments/{partnerShipmentId}/track
GET /shipments/track?trackingNumber=JD014600011234567890
""",
    )
    add_code_block(
        doc,
        """
{
  "statusCode": "in_transit",
  "statusLabel": "In transit",
  "happenedAt": "2026-06-03T12:00:00.000Z",
  "location": "Almaty Hub"
}
""",
    )
    add_table(
        doc,
        ["Response meaning", "Accepted field names"],
        [
            ["Status code", "statusCode, status, code"],
            ["Status label", "statusLabel, statusText, label"],
            ["Timestamp", "happenedAt, updatedAt, timestamp"],
            ["Location", "location, city, place"],
        ],
        [2500, 6860],
    )

    doc.add_heading("6. Generic SMS API Contract", level=1)
    doc.add_heading("6.1 Send SMS", level=2)
    add_code_block(doc, "POST /messages")
    add_code_block(
        doc,
        """
{
  "to": "+998901234567",
  "text": "Your shipment 990000000123 is ready.",
  "templateCode": "shipment_ready",
  "metadata": { "orderId": "990000000123" }
}
""",
    )
    add_code_block(
        doc,
        """
{
  "messageId": "sms_123456",
  "acceptedAt": "2026-06-03T12:00:00.000Z"
}
""",
    )
    add_callout(
        doc,
        "Required response",
        "`messageId` is required. CargoPilot also accepts `id` as an alias.",
        LIGHT_TEAL,
        TEAL,
    )
    doc.add_heading("6.2 Get delivery status", level=2)
    add_code_block(doc, "GET /messages/{messageId}")
    add_code_block(
        doc,
        """
{
  "status": "delivered",
  "deliveredAt": "2026-06-03T12:00:30.000Z"
}
""",
    )

    doc.add_heading("7. Generic Outbound Webhook Contract", level=1)
    add_body(
        doc,
        "For `webhook_sink` providers, CargoPilot sends the canonical event payload to the configured `endpointUrl`. Any `2xx` response is treated as successful.",
    )
    add_code_block(
        doc,
        """
content-type: application/json
x-cargopilot-provider-code: <provider-code>
x-cargopilot-idempotency-key: <idempotency-key>
x-cargopilot-event-type: <event-type>
authorization: Bearer <token>        # when configured
x-api-key: <api-key>                 # when configured
""",
    )
    add_table(
        doc,
        ["Result", "Classification"],
        [
            ["Any 2xx", "Successful delivery"],
            ["429", "Retryable rate limit"],
            ["Any 5xx", "Retryable provider/platform error"],
            ["Timeout, reset, DNS temporary failure, connection refused", "Retryable transport error"],
            ["Most other 4xx", "Non-retryable business/configuration error"],
            ["Invalid or missing endpoint configuration", "Non-retryable configuration error"],
        ],
        [3700, 5660],
    )

    doc.add_heading("8. Inbound Webhook Contract", level=1)
    add_body(doc, "Partners notify CargoPilot through the signed generic webhook gateway.")
    add_code_block(doc, "POST /api/integrations/webhooks/{providerCode}")
    doc.add_heading("8.1 Required signature headers", level=2)
    add_code_block(
        doc,
        """
content-type: application/json
x-signature: <hmac-signature>
x-signature-timestamp: <unix-timestamp-seconds-or-ms>
""",
    )
    add_body(doc, "Signature algorithm:")
    add_code_block(
        doc,
        """
signed_payload = "<timestamp>.<raw_body>"
signature = HMAC_SHA256_HEX(secret, signed_payload)
""",
    )
    add_callout(
        doc,
        "Replay protection",
        "CargoPilot rejects webhook timestamps outside the configured drift window. The default window is 300 seconds.",
        PALE_GOLD,
        GOLD,
    )
    doc.add_heading("8.2 Recommended webhook payload", level=2)
    add_code_block(
        doc,
        """
{
  "eventId": "evt_123456",
  "eventType": "carrier.status.updated",
  "occurredAt": "2026-06-03T12:00:00.000Z",
  "companyId": "6efc7f6d-5c31-4c5f-81b9-651ad2bd63e3",
  "aggregateType": "shipment",
  "aggregateId": "DHL-123456789",
  "payload": {
    "trackingNumber": "JD014600011234567890",
    "statusCode": "delivered",
    "statusLabel": "Delivered",
    "location": "Tashkent"
  }
}
""",
    )
    doc.add_heading("8.3 Webhook responses", level=2)
    add_table(
        doc,
        ["HTTP result", "Meaning", "Example status"],
        [
            ["202", "Signature valid and event accepted", "accepted"],
            ["200", "Duplicate event already recorded", "duplicate"],
            ["400", "Invalid signature, timestamp, provider, or payload", "rejected"],
        ],
        [1600, 5200, 2560],
    )
    add_figure(doc, diagrams["sequence"], "Figure 2. Outbound delivery sequence and persisted attempt handling.")

    doc.add_heading("9. Reliability and Operations", level=1)
    add_body(
        doc,
        "CargoPilot retries temporary failures and preserves a full attempt history. Operations teams can inspect records, force a retry, or replay failed/dead-letter records from the admin console.",
    )
    add_figure(doc, diagrams["state"], "Figure 3. Integration outbox state machine.")
    doc.add_heading("Partner HTTP behavior", level=2)
    add_bullets(
        doc,
        [
            "Return `2xx` only after accepting the operation.",
            "Return `429` when rate limited.",
            "Return `5xx` for temporary provider/platform errors.",
            "Return `409` only for a real business conflict.",
            "Do not return `2xx` for a failed business operation.",
        ],
    )
    doc.add_heading("Operational ownership", level=2)
    add_table(
        doc,
        ["CargoPilot owns", "Partner owns"],
        [
            ["Provider activation and environment selection", "API availability and contract compliance"],
            ["Credential rotation and encrypted storage", "Credential issuance and revocation"],
            ["Queue inspection, retry, and replay", "Idempotent handling of repeated requests"],
            ["Webhook verification and duplicate detection", "Correct signatures and stable event IDs"],
            ["Delivery audit trail", "Clear HTTP status codes and error responses"],
        ],
        [4680, 4680],
    )

    doc.add_heading("10. UML Data Model", level=1)
    add_figure(doc, diagrams["class"], "Figure 4. Main integration provider, secret, outbox, and attempt relationships.")
    add_table(
        doc,
        ["Model", "Business purpose"],
        [
            ["IntegrationProvider", "Defines the company-scoped external provider and runtime controls."],
            ["IntegrationProviderSecret", "Stores versioned encrypted configuration and credentials."],
            ["IntegrationOutbox", "Stores each outbound event and its delivery lifecycle."],
            ["IntegrationDeliveryAttempt", "Stores the result and diagnostics for every attempt."],
            ["IntegrationWebhookEvent", "Stores raw inbound events after signature verification."],
        ],
        [2600, 6760],
    )

    doc.add_heading("11. Testing and Production Readiness", level=1)
    add_body(doc, "Production activation requires joint acceptance of the following checks.")
    add_table(
        doc,
        ["Test", "Expected result", "Owner"],
        [
            ["Sandbox outbound request", "Partner accepts request and returns required reference", "Joint"],
            ["Idempotency replay", "Repeated key does not create duplicate business effect", "Partner"],
            ["Temporary 5xx/timeout", "CargoPilot retries and records attempts", "Joint"],
            ["Inbound signed webhook", "CargoPilot accepts valid signature", "Partner"],
            ["Invalid signature", "CargoPilot rejects webhook", "CargoPilot"],
            ["Duplicate webhook", "CargoPilot returns duplicate without reapplying effect", "Joint"],
            ["Secret rotation", "New credential works; previous credential can be revoked", "Joint"],
            ["Pause/disable provider", "No new dispatch occurs", "CargoPilot"],
            ["Production credentials", "Loaded, verified, and approved", "Joint"],
        ],
        [2650, 5000, 1710],
    )
    add_callout(
        doc,
        "Go-live gate",
        "A provider should not be activated in production until sandbox behavior, idempotency, signatures, retries, and credential rotation have been verified.",
        PALE_RED,
        RED,
    )

    doc.add_heading("Appendix A. Canonical Event Types", level=1)
    add_table(
        doc,
        ["Event type", "Typical use"],
        [
            ["order.created", "Create or notify a partner about a new order."],
            ["order.status.changed", "Notify a partner about an operational status change."],
            ["shipment.assigned", "Create/assign shipment work with a carrier."],
            ["shipment.delivered", "Notify delivery completion."],
            ["payment.intent.created", "Notify/create a payment operation."],
            ["payment.paid", "Notify confirmed payment."],
            ["support.ticket.created", "Notify a support/CRM partner."],
            ["sms.delivery.updated", "Normalize SMS delivery status."],
            ["carrier.status.updated", "Normalize carrier shipment status."],
        ],
        [3100, 6260],
    )

    doc.add_heading("Appendix B. Inbound Webhook Signing Example", level=1)
    add_code_block(
        doc,
        """
const crypto = require("crypto");

const timestamp = Math.floor(Date.now() / 1000).toString();
const rawBody = JSON.stringify(payload);
const signature = crypto
  .createHmac("sha256", webhookSecret)
  .update(`${timestamp}.${rawBody}`)
  .digest("hex");

// Send rawBody with:
// x-signature: signature
// x-signature-timestamp: timestamp
""",
    )
    add_callout(
        doc,
        "Important",
        "Generate the signature from the exact raw request body bytes sent to CargoPilot. Re-serializing JSON after signing can change whitespace or field order and invalidate the signature.",
        PALE_GOLD,
        GOLD,
    )

    doc.add_heading("Appendix C. Partner Error Response Guidance", level=1)
    add_table(
        doc,
        ["Status", "Use when", "CargoPilot behavior"],
        [
            ["200-299", "Operation accepted successfully", "Marks record sent"],
            ["400", "Request schema is invalid", "Final failure unless manually replayed after correction"],
            ["401/403", "Credential is invalid or unauthorized", "Final failure; rotate/fix credentials"],
            ["409", "Real business conflict", "Final failure; investigate business state"],
            ["429", "Partner rate limit reached", "Retries later"],
            ["500-599", "Temporary partner/platform error", "Retries later"],
        ],
        [1500, 4200, 3660],
    )

    # Apply layout and save.
    configure_sections(doc)
    for index, sec in enumerate(doc.sections):
        configure_header_footer(sec, first_page=(index == 0))
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
